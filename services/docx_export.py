"""
Converts plain-text transcripts and Markdown SOAP notes into formatted .docx files.
"""

import re
from io import BytesIO
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _set_heading_color(paragraph, hex_color: str = "6366f1") -> None:
    """Apply an RGB colour to every run in a heading paragraph."""
    r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(r, g, b)


def _add_horizontal_rule(doc: Document) -> None:
    """Insert a thin paragraph border that renders as a horizontal rule in Word."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1e2035")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _strip_inline_md(text: str) -> str:
    """Remove bold/italic markers so plain text reads cleanly in Word."""
    text = re.sub(r"\*{1,2}(.+?)\*{1,2}", r"\1", text)
    text = re.sub(r"_{1,2}(.+?)_{1,2}", r"\1", text)
    return text


def _add_body_content(doc: Document, lines: list[str]) -> None:
    """
    Parse a list of content lines and add them to the document as either
    bullet-list items or normal paragraphs.
    """
    buffer: list[str] = []

    def flush_buffer():
        if buffer:
            text = " ".join(buffer).strip()
            if text:
                p = doc.add_paragraph(_strip_inline_md(text))
                p.paragraph_format.space_after = Pt(4)
            buffer.clear()

    for line in lines:
        line = line.rstrip()
        if not line:
            flush_buffer()
            continue

        if re.match(r"^[-•*]\s", line):
            flush_buffer()
            item = _strip_inline_md(line[2:].strip())
            p = doc.add_paragraph(item, style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
        else:
            buffer.append(_strip_inline_md(line))

    flush_buffer()


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def transcript_to_docx(transcript: str) -> bytes:
    """Return a .docx byte stream for the cleaned transcript."""
    doc = Document()

    # Narrow margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.1)
        section.right_margin  = Inches(1.1)

    title = doc.add_heading("Medical Transcription", 0)
    _set_heading_color(title, "6366f1")

    meta = doc.add_paragraph(f"Date: {date.today().strftime('%B %d, %Y')}")
    meta.paragraph_format.space_after = Pt(2)
    for run in meta.runs:
        run.font.color.rgb = RGBColor(0x56, 0x58, 0x70)
        run.font.size = Pt(9)

    _add_horizontal_rule(doc)
    doc.add_paragraph()  # spacer

    for block in transcript.strip().split("\n\n"):
        block = block.strip()
        if block:
            p = doc.add_paragraph(_strip_inline_md(block))
            p.paragraph_format.space_after = Pt(6)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def soap_to_docx(soap_markdown: str) -> bytes:
    """Return a .docx byte stream for the SOAP note, with styled section headers."""
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.1)
        section.right_margin  = Inches(1.1)

    title = doc.add_heading("SOAP Note", 0)
    _set_heading_color(title, "6366f1")

    meta = doc.add_paragraph(f"Date: {date.today().strftime('%B %d, %Y')}")
    meta.paragraph_format.space_after = Pt(2)
    for run in meta.runs:
        run.font.color.rgb = RGBColor(0x56, 0x58, 0x70)
        run.font.size = Pt(9)

    _add_horizontal_rule(doc)
    doc.add_paragraph()  # spacer

    # Split on Markdown ## headers
    parts = re.split(r"^## ", soap_markdown.strip(), flags=re.MULTILINE)

    for part in parts:
        part = part.strip()
        if not part:
            continue

        lines     = part.split("\n")
        header    = lines[0].strip()
        body_lines = lines[1:]

        h = doc.add_heading(header, level=1)
        _set_heading_color(h, "818cf8")
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after  = Pt(6)

        _add_body_content(doc, body_lines)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
